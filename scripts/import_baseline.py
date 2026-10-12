"""把《菜單配方表.xlsx》轉成系統可匯入的 JSON。

用法：
    python scripts/import_baseline.py [xlsx 路徑] [輸出路徑]

預設依序在專案資料夾、private/、上一層資料夾找「菜單配方表.xlsx」，輸出到 private/baseline-seed.json
（private/ 與專案資料夾裡的 xlsx 都不進 git）。輸出後在系統「更多 → 資料匯出與匯入」選擇這個檔案匯入。

轉換規則（不猜數字：沒有寫克重或毫升的，用量一律留空，系統顯示「用量待填」）：
- 「基湯食譜」每一列 → 元件（湯底）v1。原料欄一行一項；「1 隻」「適量」這類寫法沒辦法換成重量，
  用量留空、原文寫進處理方式。用量都填齊才會自動送「試菜中」。
- 「菜單食譜」每一列 → 菜品 v1 草案。
  - 「基底湯／液體」「麵量」→ 用料
  - 「材料」是整批醬料的配方 → 另外建立元件（醬料），「醬料做法」是它的步驟；菜品引用這個醬料，每份用量留空
  - 「第一版建議作法」裡「蛋白質前置作業」底下一行一項、有克重的 → 用料；其餘文字與「前置／出餐」→ 步驟
- 試菜結果、調整方向等欄位有填的話寫進版本備註。

重新匯入時，系統裡已經有的食譜只有在配方表多了新內容（新的用量、新引用的元件、新的步驟）時才會多一個草案版本，
舊版本不會被覆蓋；配方表裡拿掉的菜不會被刪除或封存。
"""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_NAME = "菜單配方表.xlsx"
DEFAULT_OUT = ROOT / "private" / "baseline-seed.json"
SOURCE_NOTE = "由《菜單配方表》匯入"

# 配方表裡改過名字的食譜：新名稱 → 系統裡可能還在用的舊名稱。匯入時會先改名，再當成同一個食譜的新版本
RENAMED = {"奶白魚高湯": ["魚高湯（黃魚奶白）"]}

# 同一種原物料的不同寫法 → (標準名稱, 處理方式)
ALIASES = {
    "蛤蠣": ("蛤蜊", ""),
    "蔥花": ("青蔥", "切蔥花"),
    "蔥": ("青蔥", ""),
    "薑片": ("薑", "切片"),
    "生薑": ("薑", ""),
    "鹽巴": ("鹽", ""),
    "冷水": ("水", "冷水"),
    "熱水": ("水", "熱水"),
    "飲用水": ("水", "飲用水"),
    "生麵": ("生麵條", ""),
    "老菜圃汁": ("老菜圃料汁", ""),
}

CATEGORY = {
    "meat": ["去骨雞腿肉", "老母雞", "豬蹄", "後腿肉", "梅花肉"],
    "seafood": ["蛤蜊", "鱸魚", "鱸魚淨肉"],
    "produce": ["青蔥", "薑", "蒜苗"],
    "dry_goods": ["乾干貝"],
    "seasoning": [
        "鹽", "糖", "白糖", "白胡椒", "胡椒粉", "米酒", "薑汁", "大骨粉", "黃豆醬", "排骨醬", "乾黃醬",
        "甜麵醬", "十三香", "MSG", "芝麻醬", "味極鮮", "醬油", "蠔油", "陳醋",
    ],
    "oil": ["中性油", "沙拉油"],
    "grain": ["生麵條"],
}

# 店內自製的半成品：先建成原物料，備註提醒之後改成元件食譜
PREPARED = {"老菜圃料汁", "茶湯"}
VOLUME_HINTS = ("油", "汁", "水", "茶湯")
VOLUME_UNITS = {"ml", "L"}

UNIT = r"(?P<unit>kg|g|ml|L)"
NUM = r"(?P<qty>\d+(?:\.\d+)?)"
# 「100g 某某肉」
LEAD_RE = re.compile(rf"^{NUM}\s*{UNIT}\s*(?P<name>\D.*)$")
# 「某某肉 100g」；名稱裡不能有數字或頓號，避免把一整句話當成一項用料
TRAIL_RE = re.compile(rf"^(?P<name>[^\d、，。：]+?)\s*{NUM}\s*{UNIT}$")
# 「1 隻 老母雞」「適量 大骨粉」：沒辦法換成重量
COUNT_RE = re.compile(r"^(?P<count>\d+\s*[隻支根尾顆片把條塊]|適量|少許)\s*(?P<rest>.+)$")
PAREN_RE = re.compile(r"[（(](.+?)[）)]")


def clean(text) -> str:
    return re.sub(r"[ \t　]+", " ", str(text or "")).strip()


def cell_lines(cell) -> list[str]:
    return [t for t in (clean(x) for x in str(cell or "").split("\n")) if t and t != "無"]


def number(text: str):
    value = float(text)
    return int(value) if value.is_integer() else value


def parse_material(raw: str) -> dict:
    """一項用料：有克重／毫升就帶入，其餘留空並把原文留在 note"""
    notes = []
    raw = clean(raw)
    paren = PAREN_RE.search(raw)
    if paren:
        notes.append(clean(paren.group(1)))
        raw = clean(raw[: paren.start()] + raw[paren.end():])
    count = COUNT_RE.match(raw)
    if count:
        notes.insert(0, clean(count.group("count")))
        raw = clean(count.group("rest"))
    m = LEAD_RE.match(raw) or TRAIL_RE.match(raw)
    if m:
        return {"name": clean(m.group("name")), "quantity": number(m.group("qty")), "unit": m.group("unit"), "note": "；".join(notes)}
    return {"name": raw, "quantity": None, "unit": None, "note": "；".join(notes)}


def is_material_line(raw: str) -> bool:
    return parse_material(raw)["quantity"] is not None


class Registry:
    def __init__(self):
        self.items: dict[str, dict] = {}

    def resolve(self, name: str, unit: str | None) -> tuple[str, str]:
        canonical, prep = ALIASES.get(name, (name, ""))
        if canonical not in self.items:
            if unit in VOLUME_UNITS or (unit is None and canonical.endswith(VOLUME_HINTS)):
                base = "volume"
            else:
                base = "mass"
            category = next((c for c, names in CATEGORY.items() if canonical in names), "other")
            note = "自製半成品：之後建立元件食譜後，建議改為引用元件" if canonical in PREPARED else ""
            self.items[canonical] = {
                "name": canonical,
                "category": category,
                "base_dimension": base,
                "density_g_per_ml": 1 if canonical == "水" else None,
                "note": note,
            }
        return canonical, prep


def ingredient_line(item: dict, registry: Registry, group: str) -> dict:
    canonical, prep = registry.resolve(item["name"], item["unit"])
    base = registry.items[canonical]["base_dimension"]
    renamed = f"原表寫「{item['name']}」" if canonical != item["name"] and not prep else ""
    return {
        "ingredient": canonical,
        "quantity": item["quantity"],
        "unit": item["unit"] or ("ml" if base == "volume" else "g"),
        "prep_note": "；".join(p for p in [prep, renamed, item["note"]] if p),
        "group_label": group,
    }


def component_line(name: str, item: dict, group: str, note: str = "") -> dict:
    return {
        "component": name,
        "quantity": item["quantity"],
        "unit": item["unit"] or "g",
        "prep_note": "；".join(p for p in [item["note"], note] if p),
        "group_label": group,
    }


def numbered_steps(text, pattern: str) -> list[dict]:
    """「1.…⏎2.…」或「STEP1：…STEP2：…」拆成一個步驟一筆"""
    parts = [clean(p) for p in re.split(pattern, str(text or ""))]
    return [{"instruction": p} for p in parts if p and not p.endswith("：")]


def rows_as_dicts(ws):
    header = [clean(c.value) for c in ws[4]]
    for r in ws.iter_rows(min_row=5, values_only=True):
        yield dict(zip(header, r))


def version(notes: list[str], lines: list[dict], steps: list[dict], **extra) -> dict:
    return {"title": "基準版", "change_note": SOURCE_NOTE, "notes": "\n".join(notes), "lines": lines, "steps": steps, **extra}


def sauce_name(dish: str) -> str:
    base = dish.removesuffix("麵")
    return base if base.endswith("醬") else base + "醬"


def build_dishes(ws, standard_ws, registry: Registry, soup_names: list[str]) -> tuple[list[dict], list[dict]]:
    noodle_standard = {clean(r[0]): clean(r[1]) for r in standard_ws.iter_rows(min_row=5, values_only=True) if r[0]}

    def soup_for(name: str) -> str | None:
        return next((s for s in soup_names if s == name or s.endswith(name)), None)

    dishes, sauces = [], []
    category = ""
    for row in rows_as_dicts(ws):
        name = clean(row.get("原菜名"))
        if not name:
            continue
        notes = []
        if clean(row.get("分類")):
            category = clean(row["分類"])
        elif category:
            notes.append(f"分類：原表這一列沒有填，暫時沿用上一列的「{category}」，請確認。")

        lines = []
        for col, group in [("基底湯／液體", "湯底／液體"), ("麵量", "麵")]:
            for raw in cell_lines(row.get(col)):
                item = parse_material(raw)
                soup = soup_for(item["name"])
                lines.append(component_line(soup, item, group) if soup else ingredient_line(item, registry, group))

        # 「材料」是整批醬料的配方，不是一份的量
        materials = [m for m in cell_lines(row.get("材料")) if not m.endswith("：")]
        if materials:
            sauce = sauce_name(name)
            sauce_lines = [ingredient_line(parse_material(m), registry, "") for m in materials]
            sauces.append(
                {
                    "name": sauce,
                    "type": "component",
                    "component_kind": "sauce",
                    "description": f"基準版：「{name}」的整批醬料配方（第一版試作基準，需經實際試菜校正）",
                    "to_testing": all(line["quantity"] for line in sauce_lines),
                    "version": version(
                        ["批次產量與每份用量還沒有量，試做後再記錄。"],
                        sauce_lines,
                        numbered_steps(row.get("醬料做法"), r"STEP\s*\d+\s*[：:]"),
                        batch_output_unit="g",
                        serving_unit="g",
                    ),
                }
            )
            lines.append(component_line(sauce, {"quantity": None, "unit": "g", "note": ""}, "醬料", "每份用量待測"))

        steps = []
        method = str(row.get("第一版建議作法") or "")
        prep_block = cell_lines(method.split("前置：")[0])
        if prep_block:
            title = prep_block[0] if not is_material_line(prep_block[0]) else ""
            body = prep_block[1:] if title else prep_block
            # 一行一項、有克重的當作用料；其餘（舒肥溫度、整段醃料說明）留在步驟
            protein = "蛋白質" in title
            others = []
            for raw in body:
                if protein and is_material_line(raw):
                    lines.append(ingredient_line(parse_material(raw), registry, "蛋白質"))
                else:
                    others.append(raw)
            if others:
                steps.append({"instruction": chr(10).join([f"{title or '前置作業'}：", *others])})
        for label in ["前置", "出餐"]:
            m = re.search(rf"(?:^|\n)\s*{label}：(.*?)(?=\n\s*(?:前置|出餐)：|$)", method, re.S)
            content = clean(m.group(1).replace("\n", " ")) if m else ""
            if content:
                steps.append({"instruction": f"{label}：{content}"})

        std = noodle_standard.get("湯麵" if "湯" in category else "乾拌麵", "")
        for line in lines:
            if line.get("ingredient") == "生麵條" and line["quantity"] and std:
                if str(line["quantity"]) not in re.findall(r"(\d+)\s*g", std):
                    notes.append(f"麵量待確認：菜單食譜寫生麵 {line['quantity']} g，《單份標準》寫「{std}」。")
        for extra in ["試菜結果", "調整方向", "下一輪？", "建議售價", "成本／份", "SOP 編號"]:
            if row.get(extra):
                notes.append(f"{extra}：{clean(row[extra])}")

        dishes.append(
            {
                "name": name,
                "type": "dish",
                "menu_category": category,
                "description": "基準版（第一版試作基準，需經實際試菜校正）",
                "to_testing": False,
                "version": version(notes, lines, steps, serving_qty=None, serving_unit="g"),
            }
        )
    return dishes, sauces


def build_soups(ws, registry: Registry, dishes: list[dict]) -> list[dict]:
    soups = []
    for row in rows_as_dicts(ws):
        name = clean(row.get("湯底"))
        if not name:
            continue
        lines = [ingredient_line(parse_material(raw), registry, "") for raw in cell_lines(row.get("原料／階段"))]
        # 每份量：菜品引用這個湯底時最常用的量
        usages = Counter(
            line["quantity"]
            for dish in dishes
            for line in dish["version"]["lines"]
            if line.get("component") == name and line["quantity"]
        )
        serving = usages.most_common(1)[0][0] if usages else None
        notes = []
        for label in ["製成時間", "預估成品／保存", "試作觀察", "採購金額"]:
            if clean(row.get(label)):
                notes.append(f"{label}：{clean(str(row[label]).replace(chr(10), '；'))}")
        if serving:
            others = "、".join(f"{q} ml" for q in sorted(usages) if q != serving)
            notes.append(f"每份量暫以菜品最常用的 {serving} ml 計" + (f"；其他菜品用量：{others}" if others else "") + "。")
        liters = re.search(r"約\s*(\d+(?:\.\d+)?)\s*L", str(row.get("預估成品／保存") or ""))
        soups.append(
            {
                "name": name,
                "previous_names": RENAMED.get(name, []),
                "type": "component",
                "component_kind": "soup",
                "description": "基準版（第一版試作基準，需經實際試菜校正）",
                "to_testing": bool(lines) and all(line["quantity"] for line in lines),
                "version": version(
                    notes,
                    lines,
                    numbered_steps(row.get("製程"), r"(?:^|\n)\s*\d+\s*[.．、]"),
                    batch_output_qty=number(liters.group(1)) * 1000 if liters else None,
                    batch_output_unit="ml",
                    serving_qty=serving,
                    serving_unit="ml",
                ),
            }
        )
    return soups


def sheet(wb, *keywords: str):
    for ws in wb:
        if all(k in ws.title for k in keywords):
            return ws
    raise SystemExit(f"找不到名稱包含「{'」「'.join(keywords)}」的工作表；目前有：{'、'.join(wb.sheetnames)}")


def find_xlsx() -> Path:
    for folder in [ROOT, ROOT / "private", ROOT.parent]:
        if (folder / XLSX_NAME).exists():
            return folder / XLSX_NAME
    raise SystemExit(f"找不到 {XLSX_NAME}（找過專案資料夾、private/、上一層資料夾）")


def report_changes(out: Path, recipes: list[dict]) -> None:
    """和上一次產生的檔案比較，列出新增、移除與內容有變的食譜"""
    if not out.exists():
        return
    try:
        old = {r["name"]: r for r in json.loads(out.read_text(encoding="utf-8"))["recipes"]}
    except (ValueError, KeyError):
        return
    new = {r["name"]: r for r in recipes}

    def content(r: dict) -> dict:
        v = r["version"]
        return {"lines": v["lines"], "steps": v["steps"], "category": r.get("menu_category")}

    print("和上一次的檔案比較：")
    print("  新增：" + ("、".join(n for n in new if n not in old) or "無"))
    print("  不在新表裡：" + ("、".join(n for n in old if n not in new) or "無"))
    print("  內容有變：" + ("、".join(n for n in new if n in old and content(new[n]) != content(old[n])) or "無"))


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else find_xlsx()
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    registry = Registry()
    soup_ws = sheet(wb, "基", "湯", "食譜")
    soup_names = [clean(r.get("湯底")) for r in rows_as_dicts(soup_ws) if clean(r.get("湯底"))]
    dishes, sauces = build_dishes(sheet(wb, "菜單"), sheet(wb, "單份標準"), registry, soup_names)
    soups = build_soups(soup_ws, registry, dishes)
    recipes = soups + sauces + dishes
    data = {
        "source": xlsx.name,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "ingredients": sorted(registry.items.values(), key=lambda x: (x["category"], x["name"])),
        "recipes": recipes,
    }
    report_changes(out, recipes)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"輸出 {out}")
    print(f"原物料 {len(registry.items)} 項、湯底 {len(soups)} 個、醬料 {len(sauces)} 個、菜品 {len(dishes)} 道")
    pending = [r["name"] for r in soups + sauces if not r["to_testing"]]
    if pending:
        print("用量還沒填齊、不會自動送試菜的元件：" + "、".join(pending))


if __name__ == "__main__":
    main()
